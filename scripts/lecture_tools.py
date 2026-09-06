#!/usr/bin/env python3
"""lecture_tools.py - the one helper script for the lecture-script skill.

    python lecture_tools.py pptx  DECK.pptx            dump slide text, notes, diagrams
    python lecture_tools.py docx  FILE.docx            dump a Word file as plain text
    python lecture_tools.py build SCRIPT.txt [-o OUT.docx] [--wpm 130] [--max-total 20]
                                                       build the color-coded Word script

Script text format (one sentence per line):
    Title: My lecture            optional, first line only
    # Segment title              new segment (Heading 1)
    ## Sub-heading               optional sub-heading inside a segment
    Plain line                   instructor's own words            -> black
    + Line starting with plus    new writing                       -> dark red
    [Line in square brackets]    production cue, not read aloud    -> gray italic small caps
    // comment                   ignored
Blank lines are ignored.
"""
import argparse
import re
import sys
from pathlib import Path

RED = (0x8B, 0x00, 0x00)
GRAY = (0x59, 0x59, 0x59)


# ----------------------------------------------------------------- pptx dump
def dump_pptx(path):
    from pptx import Presentation
    from pptx.util import Emu  # noqa: F401  (import check)
    from lxml import etree

    prs = Presentation(path)
    ns_a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

    def shape_text(shape, out):
        if shape.shape_type == 6 and hasattr(shape, "shapes"):  # group
            for s in shape.shapes:
                shape_text(s, out)
            return
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = "".join(r.text for r in p.runs).strip()
                if t:
                    out.append(("  " * p.level) + "- " + t)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                out.append("- | " + " | ".join(cells) + " |")
        elif shape.shape_type == 13:  # picture
            out.append(f"- [image: {shape.name}]")

    for i, slide in enumerate(prs.slides, 1):
        hidden = slide._element.get("show") == "0"
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        print(f"\n=== Slide {i}{' (HIDDEN)' if hidden else ''}: {title}")
        body = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            shape_text(shape, body)
        # SmartArt / diagram text lives in separate parts python-pptx does not expose
        for rel in slide.part.rels.values():
            if "diagramData" in rel.reltype:
                root = etree.fromstring(rel.target_part.blob)
                texts = [t.text.strip() for t in root.iter(ns_a + "t") if t.text and t.text.strip()]
                if texts:
                    body.append("- [diagram] " + " / ".join(texts))
        for line in body:
            print(line)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                print("  NOTES: " + notes.replace("\n", "\n  NOTES: "))


# ----------------------------------------------------------------- docx dump
def dump_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = docx.Document(path)
    for block in doc.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(block, doc)
            text = p.text.strip()
            if not text:
                continue
            style = p.style.name if p.style is not None else ""
            m = re.match(r"Heading (\d)", style)
            prefix = "#" * int(m.group(1)) + " " if m else ("- " if "List" in style else "")
            print(prefix + text)
        elif tag == "tbl":
            t = Table(block, doc)
            for row in t.rows:
                cells = []
                for c in row.cells:
                    txt = " ".join(x.strip() for x in c.text.split("\n") if x.strip())
                    if not cells or cells[-1] != txt:  # merged cells repeat
                        cells.append(txt)
                print("| " + " | ".join(cells) + " |")
            print()


# --------------------------------------------------------------------- build
def parse_script(text):
    """Return (title, segments). A segment is dict(title, blocks); block = (kind, text)."""
    title = None
    segments = []
    cur = None
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.startswith("//"):
            continue
        if title is None and not segments and line.lower().startswith("title:"):
            title = line.split(":", 1)[1].strip()
            continue
        if line.startswith("## "):
            kind, body = "sub", line[3:].strip()
        elif line.startswith("# "):
            cur = {"title": line[2:].strip(), "blocks": []}
            segments.append(cur)
            continue
        elif line.startswith("+"):
            kind, body = "new", line[1:].strip()
        elif line.startswith("[") and line.endswith("]"):
            kind, body = "cue", line[1:-1].strip()
        else:
            kind, body = "own", line
        if cur is None:
            cur = {"title": "Segment 1", "blocks": []}
            segments.append(cur)
        cur["blocks"].append((kind, body))
    return title, segments


def words(s):
    return len(re.findall(r"[\w'’-]+", s))


def build(script_path, out_path, wpm, max_total, seg_min, seg_max):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: F401
    from docx.shared import Pt, RGBColor

    src = Path(script_path).read_text(encoding="utf-8-sig")
    title, segments = parse_script(src)
    if not segments:
        sys.exit("ERROR: no script lines found")
    title = title or Path(script_path).stem

    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)

    def run(par, text, color=None, italic=False, small_caps=False, bold=False):
        r = par.add_run(text)
        if color:
            r.font.color.rgb = RGBColor(*color)
        r.font.italic = italic
        r.font.small_caps = small_caps
        r.font.bold = bold
        return r

    doc.add_heading(title, level=0)
    key = doc.add_paragraph()
    run(key, "Color key. ", bold=True)
    run(key, "Black: the instructor's own words from the recorded class. ")
    run(key, "Dark red: new writing, not spoken in class. ", color=RED)
    run(key, "Gray: production cues, not read aloud.", color=GRAY, italic=True, small_caps=True)

    report = []
    total_own = total_new = 0
    for seg in segments:
        own = sum(words(t) for k, t in seg["blocks"] if k == "own")
        new = sum(words(t) for k, t in seg["blocks"] if k == "new")
        minutes = (own + new) / wpm
        total_own += own
        total_new += new
        doc.add_heading(seg["title"], level=1)
        meta = doc.add_paragraph()
        run(meta, f"{own + new} words · about {minutes:.1f} min at {wpm} wpm · "
                  f"{(new / (own + new) * 100) if own + new else 0:.0f}% new writing",
            color=GRAY, italic=True)
        for kind, text in seg["blocks"]:
            if kind == "sub":
                doc.add_heading(text, level=2)
                continue
            p = doc.add_paragraph()
            if kind == "own":
                run(p, text)
            elif kind == "new":
                run(p, text, color=RED)
            else:
                run(p, text, color=GRAY, italic=True, small_caps=True)
        flag = ""
        if minutes < seg_min or minutes > seg_max:
            flag = f"  <-- outside {seg_min}-{seg_max} min"
        report.append(f"  {seg['title'][:48]:<50} {own:>5} own {new:>5} new  {minutes:5.1f} min{flag}")

    doc.save(out_path)
    total = total_own + total_new
    total_min = total / wpm
    print(f"Built {out_path}")
    print(f"{'Segment':<52} {'words':>17}   time")
    print("\n".join(report))
    print(f"  {'TOTAL':<50} {total_own:>5} own {total_new:>5} new  {total_min:5.1f} min"
          f"{'  <-- over ' + str(max_total) + ' min' if total_min > max_total else ''}")
    if total:
        print(f"  New writing: {total_new / total * 100:.0f}% of spoken words")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("pptx", help="dump slide text, notes and diagram text")
    a.add_argument("path")
    b = sub.add_parser("docx", help="dump a Word document as text")
    b.add_argument("path")
    c = sub.add_parser("build", help="build the color-coded Word script")
    c.add_argument("script")
    c.add_argument("-o", "--out")
    c.add_argument("--wpm", type=int, default=130, help="speaking rate for time estimates (default 130)")
    c.add_argument("--max-total", type=float, default=20, help="flag if total exceeds this many minutes")
    c.add_argument("--seg-min", type=float, default=5)
    c.add_argument("--seg-max", type=float, default=15)
    args = ap.parse_args()

    if args.cmd == "pptx":
        dump_pptx(args.path)
    elif args.cmd == "docx":
        dump_docx(args.path)
    else:
        out = args.out or str(Path(args.script).with_suffix(".docx"))
        build(args.script, out, args.wpm, args.max_total, args.seg_min, args.seg_max)


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except AttributeError:
        pass
    main()
